import os

import numpy as np
import pandas as pd
from docx.shared import Pt

from docx_tables import add_table, new_document
from features import FACILITY_ITEMS, ITEM_LABELS, SANITATION_ITEMS, build_features, time_split
from make_figures import OUT, load, trans as tr

NAIVE_KEY = "Naive persistence"
NAIVE_LABEL = "Naïve persistence"
MODEL_DISPLAY = {"Random Forest": "Random forest"}
MINUS = "−"


def body(cells, indent=True, bold=False):
    return {"cells": cells, "indent": indent, "bold": bold}


def section(title, ncol):
    return {"cells": [title] + [""] * (ncol - 1), "section": True}


# --------------------------------------------------------------------------- #
# Table 1: training vs test set characteristics
# --------------------------------------------------------------------------- #
def table1_sample_characteristics(doc):
    df, tiers = build_features()
    common = df.dropna(subset=[f for f in tiers["Enhanced"]
                               if f not in ("region", "business", "building")])
    train, test = time_split(common)
    groups = [train, test]
    cs = pd.concat([train, test], ignore_index=True)
    mon = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

    def levels(col, ref, last=None):
        cnt = cs[col].value_counts()
        mid = [k for k in cnt.index if k != ref and k != last]
        return ([ref] if ref in cnt.index else []) + mid + ([last] if last in cnt.index else [])

    def smd_cont(a, b):
        pooled = np.sqrt((a.std() ** 2 + b.std() ** 2) / 2)
        return np.nan if pooled == 0 else (a.mean() - b.mean()) / pooled

    def smd_prop(p1, p2):
        pooled = np.sqrt((p1 * (1 - p1) + p2 * (1 - p2)) / 2)
        return np.nan if pooled == 0 else (p1 - p2) / pooled

    def smd_cell(v):
        return ["—" if not np.isfinite(v) else f"{abs(v):.2f}"]

    rows = []

    def cont_row(label, col, skewed=False):
        cells = [label]
        for g in groups:
            s = g[col]
            cells.append(f"{s.median():.2f} [{s.quantile(.25):.2f}–{s.quantile(.75):.2f}]"
                         if skewed else f"{s.mean():.1f} ± {s.std():.1f}")
        rows.append(body(cells + smd_cell(smd_cont(train[col], test[col]))))

    def prop_row(label, mask_fn):
        cells, props = [label], []
        for g in groups:
            n = int(mask_fn(g).sum())
            props.append(n / len(g))
            cells.append(f"{n:,} ({100 * n / len(g):.1f}%)")
        rows.append(body(cells + smd_cell(smd_prop(*props))))

    rows.append(section("Sample", 4))
    rows.append(body(["Transitions, <i>n</i>", f"{len(train):,}", f"{len(test):,}", ""]))
    rows.append(body(["Establishments, <i>n</i>", f"{train['cid'].nunique()}",
                      f"{test['cid'].nunique()}", ""]))
    rows.append(section("Outcomes (next period)", 4))
    cont_row("Capture rate<sup>a</sup>", "y_density", skewed=True)
    cont_row("Positive rate<sup>a</sup>", "y_posrate", skewed=True)
    rows.append(section("Monitoring and environmental predictors", 4))
    cont_row("Current capture rate<sup>a</sup>", "density", skewed=True)
    cont_row("Current positive rate<sup>a</sup>", "distribution", skewed=True)
    cont_row("Mean outdoor temperature (°C)", "temp")
    cont_row("Humidity (mean RH, %)", "humidity")
    for title, items in (("Facility audit items<sup>b</sup>", FACILITY_ITEMS),
                         ("Sanitation audit items<sup>b</sup>", SANITATION_ITEMS)):
        rows.append(section(title, 4))
        for c in items:
            prop_row(ITEM_LABELS[c], lambda g, c=c: g[c] > 0)
    rows.append(section("Business type", 4))
    for b in levels("business", "迴轉壽司店"):
        prop_row(tr(b), lambda g, b=b: g["business"] == b)
    rows.append(section("Building type", 4))
    prop_row("In-mall / department store", lambda g: g["building"] == "y")
    prop_row("Standalone / street-level", lambda g: g["building"] == "n")
    rows.append(section("Region", 4))
    for r in levels("region", "臺北市", last="Other city"):
        prop_row(tr(r), lambda g, r=r: g["region"] == r)
    rows.append(section("Calendar month", 4))
    for m in range(1, 13):
        prop_row(mon[m - 1], lambda g, m=m: g["month"] == m)

    header = [[{"t": "Characteristic", "align": "l"},
               {"t": f"Training set\n(<i>n</i> = {len(train):,})"},
               {"t": f"Test set\n(<i>n</i> = {len(test):,})"},
               {"t": "SMD"}]]
    notes = [
        "Values are <i>n</i> (%) for categorical variables, mean ± SD for continuous variables.",
        "<sup>a</sup> Median [interquartile range].",
        "<sup>b</sup> Number (%) with ≥ 1 recorded deficiency.",
        "RH, relative humidity; SD, standard deviation; SMD, absolute standardized mean "
        "difference.",
    ]
    add_table(doc, label="Table 1",
              caption=f"Characteristics of the training and held-out test sets "
                      f"(<i>N</i> = {len(cs):,} establishment-period transitions from "
                      f"{cs['cid'].nunique()} establishments)",
              header=header, rows=rows, widths=[7.4, 4.0, 4.0, 1.6],
              notes=notes, landscape=False, new_page=False)


# --------------------------------------------------------------------------- #
# Table 2: GLMM drivers of next-period activity
# --------------------------------------------------------------------------- #
SEASON_LABELS = {"Spring": "Spring", "Summer": "Summer", "Autumn": "Autumn"}
REGION_LABELS = {"Taichung": "Taichung", "NewTaipei": "New Taipei", "Nantou": "Nantou",
                 "Kaohsiung": "Kaohsiung", "HsinchuCity": "Hsinchu City",
                 "Taoyuan": "Taoyuan", "Tainan": "Tainan", "OtherCity": "Other city"}
BIZ_LABELS = {"Bakery": "Bakery", "Western": "Western-style", "HotPot": "Hot-pot",
              "Ramen": "Ramen", "Cafe": "Café", "Other": "Other",
              "Japanese": "Japanese", "Chinese": "Chinese"}
CONT_TERMS = [
    ("Current capture rate (per +0.10 captures trap<sup>−1</sup> day<sup>−1</sup>)", "density_u10"),
    ("Current positive rate (per +0.10)", "distribution_u10"),
    ("Mean outdoor temperature (per +1 °C)", "temp"),
    ("Humidity (per +1% RH)", "humidity"),
]


def table2_drivers(doc):
    cap = load("glmm_capture.csv")
    pos = load("glmm_posrate.csv")
    meta = load("glmm_meta.csv")

    def pfmt(p):
        return ("< 0.001" if p < 0.001 else f"{p:.3f}") + ("*" if p < 0.05 else "")

    def eff(e, lo, hi):
        if round(lo, 2) == 1.00 or round(hi, 2) == 1.00:
            return f"{e:.3f} ({lo:.3f}–{hi:.3f})"
        return f"{e:.2f} ({lo:.2f}–{hi:.2f})"

    def cells(df, term):
        r = df[df.term == term]
        if not len(r):
            return "—", "—", "—"
        r = r.iloc[0]
        return (f"{r.coef:+.3f}".replace("-", MINUS), eff(r.eff, r.lo, r.hi), pfmt(r.p))

    def bold_cols(term):
        cols = []
        for df, idx in ((cap, [1, 2, 3]), (pos, [4, 5, 6])):
            r = df[df.term == term]
            if len(r) and np.isfinite(r.p.values[0]) and float(r.p.values[0]) < 0.05:
                cols += idx
        return cols or False

    rows = []

    def term_row(label, term, indent):
        rows.append(body([label, *cells(cap, term), *cells(pos, term)],
                         indent=indent, bold=bold_cols(term)))

    def ref(label):
        rows.append(body([label, "0 (ref)", "1.00 (ref)", "—", "0 (ref)", "1.00 (ref)", "—"]))

    def levels(prefix, preferred):
        present = [t[len(prefix):] for t in cap["term"].astype(str)
                   if t.startswith(prefix) and len(t) > len(prefix)]
        return [k for k in preferred if k in present]

    rows.append(section("Monitoring indices and environment", 7))
    for lab, term in CONT_TERMS:
        term_row(lab, term, indent=False)
    present = set(cap["term"].astype(str))
    for title, items in (("Facility deficiencies (per +1 recorded instance)", FACILITY_ITEMS),
                         ("Sanitation deficiencies (per +1 recorded instance)", SANITATION_ITEMS)):
        rows.append(section(title, 7))
        for c in items:
            if c in present:
                term_row(ITEM_LABELS[c], c, indent=False)
    for title, ref_label, prefix, labels in (
            ("Season (ref: Winter)", "Winter", "season", SEASON_LABELS),
            ("Region (ref: Taipei)", "Taipei", "region", REGION_LABELS),
            ("Business type (ref: Conveyor-belt sushi)", "Conveyor-belt sushi", "business",
             BIZ_LABELS)):
        rows.append(section(title, 7))
        ref(ref_label)
        for key in levels(prefix, list(labels)):
            term_row(labels[key], f"{prefix}{key}", indent=True)
    rows.append(section("Building type (ref: In-mall / department store)", 7))
    ref("In-mall / department store")
    term_row("Standalone / street-level", "buildingStandalone", indent=True)

    mc = meta[meta.model.str.contains("capture")].iloc[0]
    mp = meta[meta.model.str.contains("posrate")].iloc[0]
    rows.append(section("Random effect (establishment)", 7))
    rows.append(body(["Intercept SD (<i>σ</i>)", f"{mc.sigma_u:.2f}", "—", "—",
                      f"{mp.sigma_u:.2f}", "—", "—"]))
    rows.append(body(["Intraclass correlation (ICC)", f"{mc.icc:.2f}", "—", "—",
                      f"{mp.icc:.2f}", "—", "—"]))

    header = [
        [{"t": "Predictor", "align": "l", "vmerge": True},
         {"t": "Capture rate<sup>a</sup>", "span": 3, "rule": True},
         {"t": "Positive rate<sup>b</sup>", "span": 3, "rule": True}],
        [{"t": ""},
         {"t": "<i>β</i>"}, {"t": "exp(<i>β</i>) (95% CI)"}, {"t": "<i>p</i>"},
         {"t": "<i>β</i>"}, {"t": "exp(<i>β</i>) (95% CI)"}, {"t": "<i>p</i>"}],
    ]
    notes = [
        "Bold values and * indicate <i>p</i> < 0.05 (unadjusted).",
        "<sup>a</sup> Tweedie, log link; exp(<i>β</i>), ratio of means.",
        "<sup>b</sup> Binomial, logit link; exp(<i>β</i>), odds ratio.",
        "CI, confidence interval; ref, reference category; RH, relative humidity; "
        "SD, standard deviation.",
    ]
    add_table(doc, label="Table 2",
              caption=f"Factors associated with next-period cockroach activity in generalized "
                      f"linear mixed models (<i>N</i> = {int(meta['n_obs'].max()):,} "
                      f"establishment-periods from {int(meta['n_group'].max())} establishments, "
                      f"training period)",
              header=header, rows=rows, widths=[6.9, 2.0, 4.4, 1.9, 2.0, 4.4, 1.9],
              notes=notes, landscape=True, new_page=True)


# --------------------------------------------------------------------------- #
# Table 3: regression performance
# --------------------------------------------------------------------------- #
def table3_regression(doc):
    perf = load("perf_extended.csv")
    metrics = ["r2", "rmse", "mae", "bias"]
    k = len(metrics)

    def get(target, model, fset, metric):
        r = perf[(perf.target == target) & (perf.model == model) & (perf.features == fset)]
        return float(r[metric].values[0])

    def fmt(v, metric):
        s = f"{v:+.3f}" if metric == "bias" else f"{v:.3f}"
        return s.replace("-", MINUS)

    specs = [(NAIVE_KEY, "—")] + [(m, f) for m in ("GLM", "Random Forest", "XGBoost")
                                  for f in ("Baseline", "Enhanced")]
    raw = [[get(t, m, f, met) for t in ("y_density", "y_posrate") for met in metrics]
           for m, f in specs]

    best = set()
    for j in range(2 * k):
        met = metrics[j % k]
        cand = [(i, round(raw[i][j], 3)) for i in range(1, len(raw)) if np.isfinite(raw[i][j])]
        key = ((lambda t: -t[1]) if met == "r2" else
               (lambda t: abs(t[1])) if met == "bias" else (lambda t: t[1]))
        top = key(min(cand, key=key))
        best |= {(i, j + 2) for i, v in cand if key((i, v)) == top}

    rows = []
    for i, ((m, f), vals) in enumerate(zip(specs, raw)):
        label = NAIVE_LABEL if m == NAIVE_KEY else MODEL_DISPLAY.get(m, m)
        if i >= 1 and i % 2 == 0:
            label = ""
        rows.append(body([label, f] + [fmt(v, metrics[j % k]) for j, v in enumerate(vals)],
                         indent=False, bold=[c for c in range(2 + 2 * k) if (i, c) in best]))

    sub = [{"t": "<i>R</i>²"}, {"t": "RMSE"}, {"t": "MAE"}, {"t": "Bias<sup>a</sup>"}]
    header = [
        [{"t": "Model", "align": "l", "vmerge": True},
         {"t": "Feature set", "vmerge": True},
         {"t": "Next capture rate", "span": k, "rule": True},
         {"t": "Next positive rate", "span": k, "rule": True}],
        [{"t": ""}, {"t": ""}] + sub + sub,
    ]
    pr = load("reg_preds.csv")
    n_test = int(((pr.target == "y_density") & (pr.model == NAIVE_KEY)).sum())
    notes = [
        "Bold values indicate the best performance among fitted models (naïve persistence "
        "excluded).",
        "<sup>a</sup> Mean of (predicted − observed); best, closest to 0.",
        "GLM, generalized linear model; MAE, mean absolute error; RMSE, root mean squared error.",
    ]
    add_table(doc, label="Table 3",
              caption=f"Predictive performance of Baseline and Enhanced feature sets for "
                      f"next-period cockroach activity in the held-out test set "
                      f"(<i>N</i> = {n_test:,})",
              header=header, rows=rows, widths=[2.8, 1.9] + [1.5] * 8,
              notes=notes, landscape=False, new_page=True)


# --------------------------------------------------------------------------- #
# Table 4: early-warning classification
# --------------------------------------------------------------------------- #
CLS_TASKS = [
    ("C. Established next", "Established vs not Established, next period"),
    ("D. Leaves Controlled next", "Leaving vs remaining in Controlled, next period "
                                  "(currently Controlled only)"),
]
CLS_MODELS = [(NAIVE_KEY, "—"), ("GLM", "Baseline"), ("GLM", "Enhanced"),
              ("Random Forest", "Baseline"), ("Random Forest", "Enhanced"),
              ("XGBoost", "Baseline"), ("XGBoost", "Enhanced")]
BOLD_COLS = (2, 3, 7)


def _merge_section_rows(tbl, n_head, rows):
    ncol = len(tbl.columns)
    for i, spec in enumerate(rows):
        if spec.get("section"):
            m = tbl.cell(n_head + i, 0).merge(tbl.cell(n_head + i, ncol - 1))
            for extra in m.paragraphs[1:]:
                extra._element.getparent().remove(extra._element)


def table4_classification(doc):
    cls = load("classification_metrics.csv")
    cls["features"] = cls["features"].fillna("—").astype(str)

    def num(v):
        return "—" if not np.isfinite(v) else f"{v:.3f}"

    def ci(v, lo, hi):
        return "—" if not np.isfinite(v) else f"{v:.3f} ({lo:.3f}–{hi:.3f})"

    rows = []
    for prefix, title in CLS_TASKS:
        g = cls[cls["task"].astype(str).str.startswith(prefix)]
        n, n_ev = int(g["n"].iloc[0]), int(g["TP"].iloc[0] + g["FN"].iloc[0])
        rows.append(section(f"{title}: <i>n</i> = {n:,}, events = {n_ev:,} "
                            f"({100 * n_ev / n:.1f}%)", 8))
        block = []
        for model, fset in CLS_MODELS:
            r = g[(g.model == model) & (g.features == fset)].iloc[0]
            label = NAIVE_LABEL if model == NAIVE_KEY else MODEL_DISPLAY.get(model, model)
            block.append([label, fset, num(float(r["auc"])), num(float(r["auprc"])),
                          ci(float(r["sensitivity"]), float(r["sens_lo"]), float(r["sens_hi"])),
                          ci(float(r["precision"]), float(r["prec_lo"]), float(r["prec_hi"])),
                          num(float(r["specificity"])), num(float(r["f1"]))])
        best = set()
        for j in BOLD_COLS:
            cand = [(i, float(block[i][j].split(" ")[0])) for i in range(1, len(block))
                    if block[i][j] != "—"]
            top = max(v for _, v in cand)
            best |= {(i, j) for i, v in cand if v == top}
        for i, cells in enumerate(block):
            if i >= 1 and i % 2 == 0:
                cells[0] = ""
            rows.append(body(cells, indent=False, bold=[c for c in range(8) if (i, c) in best]))

    header = [[{"t": "Model", "align": "l"}, {"t": "Feature set"}, {"t": "AUC"},
               {"t": "AUPRC"}, {"t": "Sensitivity\n(95% CI)"}, {"t": "Precision\n(95% CI)"},
               {"t": "Specificity"}, {"t": "F1"}]]
    notes = [
        "Bold values indicate the best AUC, AUPRC and F1 among fitted models (naïve "
        "persistence excluded).",
        "AUC, area under the receiver operating characteristic curve; AUPRC, area under the "
        "precision–recall curve; CI, confidence interval; GLM, generalized linear model.",
    ]
    doc.add_paragraph().paragraph_format.space_before = Pt(14)
    tbl = add_table(doc, label="Table 4",
                    caption="Early-warning classification performance for next-period "
                            "cockroach risk states in the held-out test set",
                    header=header, rows=rows, widths=[3.1, 1.7, 1.2, 1.6, 2.9, 2.9, 1.85, 1.1],
                    notes=notes, landscape=False, new_page=False)
    _merge_section_rows(tbl, len(header), rows)


def main():
    doc = new_document()
    table1_sample_characteristics(doc)
    table2_drivers(doc)
    table3_regression(doc)
    table4_classification(doc)
    out = os.path.join(OUT, "Tables_1-4.docx")
    doc.save(out)
    print(f"Tables saved to {out}")


if __name__ == "__main__":
    main()
