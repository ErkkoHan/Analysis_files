import re

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# --------------------------------------------------------------------------- #
# Table style
# --------------------------------------------------------------------------- #
FONT      = "Times New Roman"
BODY_PT   = 9.0
HEAD_PT   = 9.0
TITLE_PT  = 10.0
NOTE_PT   = 8.0

RULE_THICK = 12
RULE_THIN  = 6

_EDGES = ("top", "left", "bottom", "right")


# --------------------------------------------------------------------------- #
# Cell borders and text
# --------------------------------------------------------------------------- #
def _existing_sz(tcPr, edge):
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        return None
    el = borders.find(qn("w:" + edge))
    if el is None or el.get(qn("w:val")) in (None, "nil", "none"):
        return None
    return el.get(qn("w:sz"))


def set_rules(cell, top=None, bottom=None):
    tcPr = cell._tc.get_or_add_tcPr()
    keep_top    = _existing_sz(tcPr, "top")
    keep_bottom = _existing_sz(tcPr, "bottom")
    sz = {"top": top if top is not None else keep_top,
          "left": None,
          "bottom": bottom if bottom is not None else keep_bottom,
          "right": None}

    old = tcPr.find(qn("w:tcBorders"))
    if old is not None:
        tcPr.remove(old)
    borders = OxmlElement("w:tcBorders")
    for edge in _EDGES:
        e = OxmlElement("w:" + edge)
        if sz[edge]:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(sz[edge]))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), "000000")
        else:
            e.set(qn("w:val"), "nil")
        borders.append(e)
    tcPr.insert_element_before(
        borders, "w:shd", "w:noWrap", "w:tcMar", "w:textDirection",
        "w:tcFitText", "w:vAlign", "w:hideMark")


def _runs(par, text, bold=False, size=BODY_PT, italic=False):
    for chunk in re.split(r"(<i>.*?</i>|<sup>.*?</sup>)", str(text)):
        if chunk == "":
            continue
        ital, sup = italic, False
        if chunk.startswith("<i>") and chunk.endswith("</i>"):
            chunk, ital = chunk[3:-4], True
        elif chunk.startswith("<sup>") and chunk.endswith("</sup>"):
            chunk, sup = chunk[5:-6], True
        for k, seg in enumerate(chunk.split("\n")):
            if k:
                par.add_run().add_break()
            if seg == "":
                continue
            r = par.add_run(seg)
            r.bold, r.italic = bold, ital
            r.font.size = Pt(size)
            r.font.name = FONT
            r.font.superscript = sup
            r.font.color.rgb = RGBColor(0, 0, 0)
            r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


_ALIGN = {"l": WD_ALIGN_PARAGRAPH.LEFT,
          "c": WD_ALIGN_PARAGRAPH.CENTER,
          "r": WD_ALIGN_PARAGRAPH.RIGHT}


def fill(cell, text, *, bold=False, align="c", size=BODY_PT, indent_cm=0.0):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = _ALIGN[align]
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(1.5), Pt(1.5)
    pf.line_spacing = 1.0
    pf.left_indent = Cm(indent_cm)
    _runs(p, text, bold=bold, size=size)


# --------------------------------------------------------------------------- #
# Table layout
# --------------------------------------------------------------------------- #
def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def _fix_grid(tbl, widths):
    grid = tbl._tbl.find(qn("w:tblGrid"))
    cols = grid.findall(qn("w:gridCol"))
    for gc, wcm in zip(cols, widths):
        gc.set(qn("w:w"), str(int(Cm(wcm).twips)))
    tblPr = tbl._tbl.tblPr
    old = tblPr.find(qn("w:tblW"))
    if old is not None:
        tblPr.remove(old)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(int(Cm(sum(widths)).twips)))
    w.set(qn("w:type"), "dxa")
    tblPr.insert_element_before(
        w, "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
        "w:tblLayout", "w:tblCellMar", "w:tblLook")


def _page(sec, landscape=False, margin_cm=2.0):
    w, h = Cm(21.0), Cm(29.7)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = h, w
    else:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = w, h
    for a in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, a, Cm(margin_cm))


# --------------------------------------------------------------------------- #
# Document and table builder
# --------------------------------------------------------------------------- #
def new_document(margin_cm=2.0):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(BODY_PT)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st.paragraph_format
    pf.space_before, pf.space_after = Pt(0), Pt(0)
    pf.line_spacing = 1.0
    _page(doc.sections[0], landscape=False, margin_cm=margin_cm)
    return doc


def add_table(doc, *, label, caption, header, rows, widths,
              notes=(), landscape=False, new_page=True, margin_cm=2.0):
    sec = doc.add_section(WD_SECTION.NEW_PAGE) if new_page else doc.sections[-1]
    _page(sec, landscape=landscape, margin_cm=margin_cm)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    _runs(p, label, bold=True, size=TITLE_PT)
    _runs(p, " " + caption, bold=False, size=TITLE_PT)

    ncol = len(widths)
    tbl = doc.add_table(rows=0, cols=ncol)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl._tbl.tblPr.append(layout)
    _fix_grid(tbl, widths)

    def _new_row():
        r = tbl.add_row()
        for c, wcm in enumerate(widths):
            r.cells[c].width = Cm(wcm)
        return r

    nhead = len(header)
    head_rows = [_new_row() for _ in range(nhead)]
    spanner_rules = []
    vmerges = []
    for i, spec in enumerate(header):
        col = 0
        for item in spec:
            span = int(item.get("span", 1))
            a, b = col, col + span - 1
            cell = head_rows[i].cells[a]
            if span > 1:
                cell = cell.merge(head_rows[i].cells[b])
            fill(cell, item.get("t", ""), bold=True, size=HEAD_PT,
                 align=item.get("align", "c"))
            if item.get("rule") and span > 1:
                spanner_rules.append((i, a, b))
            if item.get("vmerge") and i + 1 < nhead:
                vmerges.append((i, a, item))
            col += span
    for (i, a, item) in vmerges:
        merged = tbl.cell(i, a).merge(tbl.cell(i + 1, a))
        fill(merged, item.get("t", ""), bold=True, size=HEAD_PT,
             align=item.get("align", "c"))
    for r in head_rows:
        for c in r.cells:
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for r in head_rows:
        _repeat_header(r)
        _cant_split(r)

    body_rows = []
    for spec in rows:
        r = _new_row()
        _cant_split(r)
        body_rows.append(r)
        cells = spec["cells"]
        is_sec = bool(spec.get("section"))
        bold = spec.get("bold", False)
        for c in range(ncol):
            txt = cells[c] if c < len(cells) else ""
            b = (bold is True) or (isinstance(bold, (list, tuple, set)) and c in bold) or is_sec
            fill(r.cells[c], txt,
                 bold=b,
                 align="l" if c == 0 else "c",
                 indent_cm=0.30 if (c == 0 and spec.get("indent") and not is_sec) else 0.0)

    for c in range(ncol):
        set_rules(tbl.cell(0, c), top=RULE_THICK)
    for (i, a, b) in spanner_rules:
        for c in range(a, b + 1):
            set_rules(tbl.cell(i, c), bottom=RULE_THIN)
    for c in range(ncol):
        set_rules(tbl.cell(nhead - 1, c), bottom=RULE_THIN)
    last = len(tbl.rows) - 1
    for c in range(ncol):
        set_rules(tbl.cell(last, c), bottom=RULE_THICK)

    for k, note in enumerate(notes):
        if not note:
            continue
        np_ = doc.add_paragraph()
        npf = np_.paragraph_format
        npf.space_before = Pt(5 if k == 0 else 2)
        npf.space_after = Pt(0)
        npf.line_spacing = 1.0
        _runs(np_, note, size=NOTE_PT)
    return tbl
